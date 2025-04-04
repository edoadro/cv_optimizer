import json
from docx import Document
import sys
import os
import re

# project specific packages
sys.path.append(os.path.join(os.path.dirname(__file__), 'src'))
from llm_interface import generate_content
from replace_utils import replace_placeholders_in_doc

def load_data():
    with open("data/profile.json") as f:
        user_data = json.load(f)

    with open("data/job_description.txt") as f:
        job_description = f.read()

    return user_data, job_description

def main():
    user_data, job_description = load_data()
    doc = Document("templates/cv_template.docx")

    replace_placeholders_in_doc(doc, user_data, job_description)

    output_path = "output/cv_final.docx"
    doc.save(output_path)
    print(f"✅ CV generated and saved to {output_path}")

if __name__ == "__main__":
    main()